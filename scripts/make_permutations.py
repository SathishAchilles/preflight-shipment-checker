"""Every scenario as one test workbook in the judges' format, with the expected result and a chat script for each booking.

Usage: python3 make_permutations.py <judges template .xlsx>   (writes ../assets/CargoPass_All_Permutations.xlsx)

Groups
  RC  every combination of the ten rules: 2^8 pass/fail x R-06 (pass, check, fail) x R-08 (pass, fail, not needed) = 2,304
  CM  carrier choice: 4 priorities x 3 countries x 1-4 pallets x 8 weights per pallet around each carrier limit = 384
  RT  every route: 4 collection cities x 7 destinations x warehouse / business address = 56
  EC  edge cases: rule boundaries, HS code formats, messy values, dangerous goods, multi-item bookings

Expected results come from the rulebook and the rate card in the template (sheets 3 and 4), not from the app.
Ready dates are 05-03-2027 (R-10 cases use 01-10-2026); regenerate with later dates after March 2027.
"""
import datetime, sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

wb = load_workbook(sys.argv[1] if len(sys.argv) > 1 else 'Case2_Shipment_Booking_Test_Data_Format.xlsx')
B, I = wb['Bookings'], wb['Items']
READY, PAST = '05-03-2027', '01-10-2026'
CITY = {'Germany': ['Berlin', 'Frankfurt', 'Hamburg'], 'France': ['Paris', 'Lyon'], 'Netherlands': ['Amsterdam', 'Rotterdam']}
COUNTRY = {c: k for k, v in CITY.items() for c in v}
ORIGINS = ['London', 'Manchester', 'Birmingham', 'Leeds']

# Rate card, read from the template so the expected carrier follows the judges' sheet.
CARRIERS = []
for row in wb['Carrier options'].iter_rows(min_row=3, values_only=True):
    if row[0]:
        CARRIERS.append({'carrier': row[0], 'dest': row[3], 'price': row[4], 'transit': row[6], 'on_time': row[7], 'max_w': row[8]})


def pick(country, priority, pallets):
    """The carrier each priority should choose (Balanced: start cheapest, step faster only if % time saved >= % cost added)."""
    opts = [dict(c, cost=c['price'] * pallets) for c in CARRIERS if c['dest'] == country]
    if priority == 'Cheapest':
        return min(opts, key=lambda c: (c['cost'], c['transit'], -c['on_time']))
    if priority == 'Fastest':
        return min(opts, key=lambda c: (c['transit'], -c['on_time'], c['cost']))
    if priority == 'Most reliable':
        return min(opts, key=lambda c: (-c['on_time'], c['transit'], c['cost']))
    opts.sort(key=lambda c: (c['cost'], c['transit']))
    p = opts[0]
    for nx in opts[1:]:
        if nx['transit'] < p['transit'] and (nx['cost'] - p['cost']) / p['cost'] <= (p['transit'] - nx['transit']) / p['transit'] + 1e-9:
            p = nx
    return p


bookings, items, expected = [], [], []
GOODS = {'mug': ('Ceramic coffee mug, 350 ml', '691200', 'ceramic coffee mugs'),
         'powerbank': ('Lithium-ion power bank 10,000 mAh', '850760', 'lithium-ion power banks'),
         'speaker': ('Portable Bluetooth speaker', '851822', 'portable Bluetooth speakers')}


def add(bid, group, scenario, status, breaks, b, its, chat, carrier=''):
    bookings.append([bid, b.get('seller', 'Northfield Goods'), b.get('origin', 'Manchester'), b.get('origin_country', 'United Kingdom'), b.get('city', 'Frankfurt'),
                     b.get('country', 'Germany'), b.get('dtype', 'Warehouse'), b.get('ready', READY), b.get('priority', 'Cheapest'), b.get('pallets', 3),
                     b.get('cartons', 90), b.get('declared', 5400), b.get('currency', 'GBP'), b.get('gross', 620), b.get('appt', 'Yes')])
    for it in its:
        items.append([bid, it.get('sku', 'MUG-01'), it.get('desc', GOODS['mug'][0]), it.get('hs', '691200'), it.get('qty', 1080), it.get('price', 5),
                      it.get('upc', 12), it.get('cartons', 90), it.get('uw', 0.5), it.get('coo', 'United Kingdom')])
    expected.append([bid, group, scenario, status, ', '.join(breaks) or 'none', carrier] + (chat + ['', '', ''])[:3])


def chat_text(b, it, noun, extra=''):
    """One message a seller would type for this booking."""
    dest = f"a {b.get('city', 'Frankfurt')} warehouse" if b.get('dtype', 'Warehouse') == 'Warehouse' else f"a business address in {b.get('city', 'Frankfurt')}"
    parts = [f"Hi, this is {b.get('seller', 'Northfield Goods')}. Ship {b.get('pallets', 3)} pallet{'s' if b.get('pallets', 3) != 1 else ''} from {b.get('origin', 'Manchester')} to {dest}, ready {b.get('ready', READY)}."]
    line = f"{it.get('qty', 1080)} {noun} at £{it.get('price', 5)} each, {it.get('upc', 12)} per carton, {it.get('uw', 0.5)} kg each"
    if it.get('hs', '691200'):
        line += f", HS {it.get('hs', '691200')}"
    if it.get('coo', 'United Kingdom'):
        line += ', made in the UK'
    parts.append(line + '.')
    parts.append(f"Total gross weight {fmt(b.get('gross', 620))} kg.")
    if b.get('dtype', 'Warehouse') == 'Warehouse':
        parts.append('Delivery slot booked.' if b.get('appt', 'Yes') == 'Yes' else 'No delivery slot booked yet.')
    parts.append(b.get('priority', 'Cheapest') + '.')
    return ' '.join(parts) + extra


def fmt(v):
    return str(int(v)) if float(v).is_integer() else str(v)


# ---- RC: every combination of the ten rules
RED, BIN = ['R-01', 'R-02', 'R-03', 'R-04', 'R-05'], ['R-01', 'R-02', 'R-03', 'R-04', 'R-05', 'R-07', 'R-09', 'R-10']
n = 0
for mask in range(256):
    for r6 in ('pass', 'check', 'fail'):
        for r8 in ('pass', 'fail', 'na'):
            n += 1
            on = {k: bool(mask & (1 << j)) for j, k in enumerate(BIN)}
            b, it = {}, {}
            goods = {'pass': 'mug', 'check': 'speaker', 'fail': 'powerbank'}[r6]
            it['desc'], it['hs'], noun = GOODS[goods]
            if on['R-01']: it['hs'] = ''
            if on['R-02']: b['declared'] = 5940
            if on['R-03']: it['upc'] = 10
            if on['R-04']: b['cartons'] = 95
            if on['R-05']: b['gross'] = 1650
            if on['R-07']: it['uw'] = round(b.get('gross', 620) / 1080 * 1.5, 3)
            if r8 == 'fail': b['appt'] = 'No'
            if r8 == 'na': b['dtype'], b['appt'] = 'Business address', 'No'
            if on['R-09']: it['coo'] = ''
            if on['R-10']: b['ready'] = PAST
            breaks = [k for k in BIN if on[k]] + (['R-06'] if r6 == 'fail' else ['R-06 (check)'] if r6 == 'check' else []) + (['R-08'] if r8 == 'fail' else [])
            breaks.sort(key=lambda s: s[:4])
            reds = sum(on[k] for k in RED) + (r6 == 'fail')
            ambers = sum(on[k] for k in ('R-07', 'R-09', 'R-10')) + (r6 == 'check') + (r8 == 'fail')
            status = 'Red' if reds else 'Amber' if ambers else 'Green'
            extra = (' Total declared value £5,940.' if on['R-02'] else '') + (f' The {noun} come in 90 cartons.' if on['R-03'] else '') + (' 95 cartons in total.' if on['R-04'] else '')
            follow = (["don't know"] if on['R-01'] else []) + (["don't know where they were made"] if on['R-09'] else [])
            add(f'RC-{n:04d}', 'Rule combinations', f"{bin(mask)[2:].zfill(8)} · R-06 {r6} · R-08 {r8}", status, breaks, b, [it], [chat_text(b, it, noun, extra)] + follow)

# ---- CM: carrier choice around every weight limit
n = 0
for country, city in (('Germany', 'Frankfurt'), ('France', 'Paris'), ('Netherlands', 'Amsterdam')):
    for pr in ('Cheapest', 'Fastest', 'Most reliable', 'Balanced'):
        for pallets in (1, 2, 3, 4):
            for wpp in (300, 400, 400.5, 500, 500.5, 750, 750.5, 800):
                n += 1
                c = pick(country, pr, pallets)
                gross = round(wpp * pallets, 1)
                b = {'city': city, 'country': country, 'priority': pr, 'pallets': pallets, 'gross': gross}
                it = {'uw': 0.1}
                fits = wpp <= c['max_w']
                add(f'CM-{n:03d}', 'Carrier matrix', f'{country} · {pr} · {pallets} pallet(s) · {wpp} kg per pallet', 'Green' if fits else 'Red', [] if fits else ['R-05'],
                    b, [it], [chat_text(b, it, 'ceramic coffee mugs')], f"{c['carrier']} · £{c['cost']:,.0f}")

# ---- RT: every route, warehouse and business address
n = 0
for origin in ORIGINS:
    for city, country in COUNTRY.items():
        for dtype in ('Warehouse', 'Business address'):
            n += 1
            b = {'origin': origin, 'city': city, 'country': country, 'dtype': dtype, 'appt': 'Yes' if dtype == 'Warehouse' else 'No'}
            add(f'RT-{n:03d}', 'Routes', f'{origin} → {city} · {dtype}', 'Green', [], b, [{}], [chat_text(b, {}, 'ceramic coffee mugs')])

# ---- EC: edge cases (chat script only where a seller could say it)
EC = []


def ec(scenario, status, breaks, b=None, its=None, chat=None):
    EC.append((scenario, status, breaks, b or {}, its or [{}], chat))


def said(b, it=None, noun='ceramic coffee mugs', extra=''):
    return [chat_text(b, it or {}, noun, extra)]


ec('Declared value exactly +1% passes', 'Green', [], {'declared': 5454}, chat=said({}, extra=' Total declared value £5,454.'))
ec('Declared value just over +1% fails', 'Red', ['R-02'], {'declared': 5454.01}, chat=said({}, extra=' Total declared value £5,454.01.'))
ec('Declared value exactly −1% passes', 'Green', [], {'declared': 5346}, chat=said({}, extra=' Total declared value £5,346.'))
ec('Exactly 500 kg per pallet passes (Carrier A)', 'Green', [], {'gross': 1500}, chat=said({'gross': 1500}))
ec('500.1 kg per pallet fails (Carrier A)', 'Red', ['R-05'], {'gross': 1500.3}, chat=said({'gross': 1500.3}))
ec('Gross weight equal to the goods passes', 'Green', [], {'gross': 540}, chat=said({'gross': 540}))
ec('Gross weight 0.1 kg under the goods', 'Amber', ['R-07'], {'gross': 539.9}, chat=said({'gross': 539.9}))
ec('Quantity one unit off the packing', 'Red', ['R-03'], {'declared': 5405}, [{'qty': 1081}], chat=said({}, {'qty': 1081}, extra=' The ceramic coffee mugs come in 90 cartons.'))
ec('HS code with a dot (6912.00)', 'Green', [], its=[{'hs': '6912.00'}], chat=said({}, {'hs': '6912.00'}))
ec('HS code with spaces ("HS 6912 00")', 'Green', [], its=[{'hs': 'HS 6912 00'}])
ec('HS code with a letter (69A200)', 'Red', ['R-01'], its=[{'hs': '69A200'}])
ec('HS code of 4 digits (6912)', 'Red', ['R-01'], its=[{'hs': '6912'}], chat=said({}, {'hs': '6912'}))
ec('HS code of 12 digits passes (advisory only)', 'Green', [], its=[{'hs': '691200100099'}])
ec('HS code not in the UK tariff passes R-01 (advisory only)', 'Green', [], its=[{'hs': '999999'}], chat=said({}, {'hs': '999999'}))
ec('HS code stored as a number lost its leading zero (90111)', 'Red', ['R-01'], its=[{'desc': 'Coffee beans, not roasted, 1 kg bag', 'hs': 90111, 'uw': 0.4}])
ec('Ready date as an Excel date number', 'Green', [], {'ready': datetime.datetime(2027, 3, 5)})
ec('Ready date as "5 Mar 2027"', 'Green', [], {'ready': '5 Mar 2027'}, chat=said({'ready': '5 Mar 2027'}))
ec('Impossible date 31-02-2027', 'Red', [], {'ready': '31-02-2027'})
ec('Ready date missing', 'Red', [], {'ready': ''})
ec('Declared value written "£5,400"', 'Green', [], {'declared': '£5,400'})
ec('Gross weight written "620 kg"', 'Green', [], {'gross': '620 kg'})
ec('Pallets not a number ("abc")', 'Red', [], {'pallets': 'abc'})
ec('Zero pallets', 'Red', [], {'pallets': 0})
ec('Negative declared value', 'Red', [], {'declared': -5})
ec('Origin country written "UK"', 'Green', [], {'origin_country': 'UK'})
ec('Destination "frankfurt am main", country "DE"', 'Green', [], {'city': 'frankfurt am main', 'country': 'DE'})
ec('Country of origin written "PRC"', 'Green', [], its=[{'coo': 'PRC'}])
ec('Paris with Germany as the country', 'Red', [], {'city': 'Paris', 'country': 'Germany'})
ec('Destination with no route (Madrid, Spain)', 'Red', [], {'city': 'Madrid', 'country': 'Spain'})
ec('Collection city outside the four hubs (Glasgow)', 'Red', [], {'origin': 'Glasgow'})
ec('Destination type "Amazon FC"', 'Green', [], {'dtype': 'Amazon FC'})
ec('Destination type "shop", no slot', 'Green', [], {'dtype': 'shop', 'appt': 'No'})
ec('Priority "asap"', 'Green', [], {'priority': 'asap'})
ec('Priority blank (Balanced is used)', 'Amber', [], {'priority': ''})
ec('Appointment written "booked"', 'Green', [], {'appt': 'booked'})
ec('Currency blank (GBP is used)', 'Green', [], {'currency': ''})
ec('Item description missing', 'Red', [], its=[{'desc': ''}])
for desc, status, breaks, noun in (('Methanol 1 L', 'Red', ['R-06'], 'bottles of methanol, 1 litre'), ('Lithium e-bike', 'Red', ['R-06'], 'lithium e-bikes'),
                                    ('Aerosol dry shampoo', 'Red', ['R-06'], 'aerosol dry shampoo cans'), ('Perfume 50 ml', 'Red', ['R-06'], 'bottles of perfume, 50 ml'),
                                    ('Spare batteries not included', 'Green', [], None), ('Pool chlorine tablets 5 kg', 'Green', [], None),
                                    ('Party poppers, box of 50', 'Green', [], 'boxes of party poppers'), ('12V car battery', 'Amber', ['R-06 (check)'], None),
                                    ('Spirit level 60 cm', 'Green', [], 'spirit levels, 60 cm'), ('Matchbox car set', 'Green', [], 'matchbox car sets'),
                                    ('Charcoal grey wool scarf', 'Green', [], 'charcoal grey wool scarves'), ('Chlorine-free pool float', 'Green', [], None),
                                    ('Fridge magnets, set of 6', 'Green', [], None), ('Soy wax candle', 'Green', [], 'soy wax candles')):
    ec(f'Goods: "{desc}"', status, breaks, its=[{'desc': desc}], chat=said({}, noun=noun) if noun else None)
two = [{}, {'sku': 'CAN-01', 'desc': 'Soy wax candle', 'hs': '340600', 'qty': 200, 'price': 4, 'upc': 10, 'cartons': 20, 'uw': 0.3}]
ec('Two items, declared value covers only the first', 'Red', ['R-02'], {'cartons': 110}, two)
ec('Two items, total cartons count only the first', 'Red', ['R-04'], {'declared': 6200}, two)
ec('Two items, everything adds up', 'Green', [], {'cartons': 110, 'declared': 6200, 'gross': 720}, two)
for k, (scenario, status, breaks, b, its, chat) in enumerate(EC, 1):
    add(f'EC-{k:03d}', 'Edge cases', scenario, status, breaks, b, its, chat or ['(file only)'])

# ---- Write the workbook: the template's own sheets, its example row cleared, plus Expected results
for ws in (B, I):
    for c in range(1, ws.max_column + 1):
        ws.cell(3, c).value = None
for r, row in enumerate(bookings, start=3):
    for c, v in enumerate(row, start=1):
        cell = B.cell(r, c, v)
        if isinstance(v, datetime.datetime):
            cell.number_format = 'DD-MM-YYYY'
for r, row in enumerate(items, start=3):
    for c, v in enumerate(row, start=1):
        cell = I.cell(r, c, v)
        if c == 4 and isinstance(v, str):
            cell.number_format = '@'
E = wb.create_sheet('Expected results', 2)
E['A1'] = (f'All scenarios — {len(bookings)} bookings. Expected results follow the rulebook and the rate card (sheets "Rulebook" and "Carrier options"). '
           'Chat script: type message 1, then answer the follow-ups in order.')
E['A1'].font = Font(name='Arial', bold=True, size=11)
hdr = ['Booking ID', 'Group', 'Scenario', 'Expected status', 'Rules expected to break', 'Expected carrier', 'Chat message 1', 'Chat follow-up 2', 'Chat follow-up 3']
for c, h in enumerate(hdr, 1):
    x = E.cell(2, c, h)
    x.font = Font(name='Arial', bold=True, color='FFFFFF')
    x.fill = PatternFill('solid', fgColor='1F3864')
    x.alignment = Alignment(vertical='center', wrap_text=True)
fills = {'Red': 'F8D7D3', 'Amber': 'FCEBC7', 'Green': 'D6EFDF'}
thin = Side(style='thin', color='C9D2DA')
for r, row in enumerate(expected, 3):
    for c, v in enumerate(row, 1):
        x = E.cell(r, c, v)
        x.font = Font(name='Arial', size=10)
        x.alignment = Alignment(vertical='top', wrap_text=c in (3, 7))
        x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    E.cell(r, 4).fill = PatternFill('solid', fgColor=fills[row[3]])
for c, w in enumerate([11, 17, 44, 13, 26, 18, 90, 26, 30], 1):
    E.column_dimensions[get_column_letter(c)].width = w
E.freeze_panes = 'A3'
E.auto_filter.ref = f'A2:I{len(expected) + 2}'
wb.save('../assets/CargoPass_All_Permutations.xlsx')
tally = {s: sum(1 for e in expected if e[3] == s) for s in ('Red', 'Amber', 'Green')}
print(len(bookings), 'bookings,', len(items), 'item rows,', tally)
