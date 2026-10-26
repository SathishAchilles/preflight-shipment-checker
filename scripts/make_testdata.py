import datetime, copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = load_workbook('Case2_Shipment_Booking_Test_Data_Format.xlsx')
B, I = wb['Bookings'], wb['Items']

def bk(id, seller, oc, dc, dcty, dtype, date, prio, pal, ctn, val, gross, appt, cur='GBP'):
    return [id, seller, oc, 'United Kingdom', dc, dcty, dtype, date, prio, pal, ctn, val, cur, gross, appt]
def it(id, sku, desc, hs, q, p, upc, c, uw, coo):
    return [id, sku, desc, hs, q, p, upc, c, uw, coo]

bookings, items, expected = [], [], []
def add(b, its, planted, status, rules, carrier):
    bookings.append(b); items.extend(its); expected.append([b[0], planted, status, rules, carrier])

add(bk('TB-01','Kettle & Co','London','Berlin','Germany','Warehouse','18-11-2026','Cheapest',2,40,3600,520,'Yes'),
    [it('TB-01','TPT-01','Stoneware teapot, 1 litre',691200,240,15,6,40,1.6,'United Kingdom')],
    'Nothing — clean control booking','Green','—','Carrier A · Road economy (£360)')
add(bk('TB-02','Harbour Textiles','Leeds','Paris','France','Business address','19-11-2026','Balanced',1,30,3900,200,'No'),
    [it('TB-02','TSH-01','Organic cotton T-shirt',None,400,6,20,20,0.2,'Turkey'),
     it('TB-02','SCK-01','Cotton socks, 3-pack',6115,500,3,50,10,0.1,'India')],
    'TSH-01 has no HS code; SCK-01 has only 4 digits','Red','R-01','Carrier B · Road standard (£210)')
add(bk('TB-03','Lumen Candles','Birmingham','Amsterdam','Netherlands','Warehouse','20-11-2026','Fastest',2,70,3000,520,'Yes'),
    [it('TB-03','CAN-01','Soy wax candle in glass jar',340600,600,5,12,50,0.45,'United Kingdom'),
     it('TB-03','CAN-02','Tealight candles, 50-pack',340600,200,4,10,20,0.6,'China')],
    'Declared £3,000 but lines total £3,800 — CAN-02 (£800) left out','Red','R-02','Carrier C · Road express (£640)')
add(bk('TB-04','Bramble Books','Manchester','Lyon','France','Business address','17-11-2026','Cheapest',1,25,2500,310,'No'),
    [it('TB-04','NB-01','A5 lined notebook',482020,1000,2.5,48,25,0.25,'United Kingdom')],
    '48 per carton × 25 cartons = 1,200, but quantity is 1,000','Red','R-03','Carrier A · Road economy (£170)')
add(bk('TB-05','Alder Home','London','Hamburg','Germany','Warehouse','23-11-2026','Most reliable',2,60,5500,640,'Yes'),
    [it('TB-05','CUS-01','Feather cushion, 45 cm',940490,300,9,10,30,0.8,'India'),
     it('TB-05','TWL-01','Cotton bath towel',630260,400,7,10,40,0.6,'Turkey')],
    'Booking says 60 cartons; items total 70','Red','R-04','Carrier C · Road express (£680)')
add(bk('TB-06','Ironside Cookware','Birmingham','Frankfurt','Germany','Warehouse','18-11-2026','Cheapest',3,60,7200,1650,'Yes'),
    [it('TB-06','CAS-01','Enamelled cast iron casserole dish, 24 cm',732392,240,30,4,60,6.2,'China')],
    '550 kg per pallet > Carrier A limit of 500 kg (cheapest choice)','Red','R-05 (fix: switch to Carrier B, £660)','Carrier A refuses → Carrier B · Road standard (£660)')
add(bk('TB-07','Granite & Grain','Leeds','Rotterdam','Netherlands','Business address','17-11-2026','Fastest',2,25,2100,900,'No'),
    [it('TB-07','PM-01','Granite pestle and mortar',680293,150,14,6,25,5.2,'Vietnam')],
    '450 kg per pallet > Carrier C limit of 400 kg (fastest choice)','Red','R-05 (fix: re-pack onto 3 pallets for C, or switch to B)','Carrier C after re-pack (£960) or Carrier B (£400)')
add(bk('TB-08','VoltBox','London','Berlin','Germany','Warehouse','19-11-2026','Cheapest',1,20,3300,180,'Yes'),
    [it('TB-08','PB-01','Lithium-ion power bank 10,000 mAh',850760,200,12,20,10,0.25,'China'),
     it('TB-08','DSH-02','Aerosol dry shampoo 200 ml',330590,300,3,30,10,0.25,'United Kingdom')],
    'Lithium battery and aerosol items','Red','R-06','Carrier A · Road economy (£180)')
add(bk('TB-09','Pebble Audio','Manchester','Paris','France','Warehouse','20-11-2026','Balanced',1,10,2500,120,'Yes'),
    [it('TB-09','SPK-01','Portable Bluetooth speaker',851822,100,25,10,10,0.9,'China')],
    'Bluetooth speaker probably holds a lithium battery — needs confirming','Amber','R-06 (needs confirmation)','Carrier B · Road standard (£210)')
add(bk('TB-10','Oak & Ember','Leeds','Amsterdam','Netherlands','Business address','18-11-2026','Cheapest',2,30,3600,600,'No'),
    [it('TB-10','CB-01','Oak chopping board, large',441990,300,12,10,30,2.2,'United Kingdom')],
    'Gross 600 kg < goods 660 kg','Amber','R-07','Carrier A · Road economy (£320)')
add(bk('TB-11','Marlow Kitchen','London','Frankfurt','Germany','Warehouse','17-11-2026','Fastest',1,10,1920,130,'No'),
    [it('TB-11','WB-01','Stainless steel water bottle, 750 ml',732393,240,8,24,10,0.35,'China')],
    'Warehouse delivery with no appointment','Amber','R-08','Carrier C · Road express (£340)')
add(bk('TB-12','Saffron Lane','Birmingham','Lyon','France','Business address','19-11-2026','Most reliable',1,18,2960,140,'No'),
    [it('TB-12','SPT-01','Brass spice tin set',741810,120,18,12,10,0.7,None),
     it('TB-12','TT-02','Cotton tea towel',630291,200,4,25,8,0.08,'India')],
    'SPT-01 has no country of origin','Amber','R-09','Carrier C · Road express (£330)')
add(bk('TB-13','Fern & Fable','Manchester','Hamburg','Germany','Warehouse','15-10-2026','Balanced',1,15,2700,230,'Yes'),
    [it('TB-13','PZ-01','Wooden jigsaw puzzle, 500 pieces',950300,300,9,20,15,0.6,'United Kingdom')],
    'Ready date 15-10-2026 is in the past','Amber','R-10','Carrier B · Road standard (£220)')
add(bk('TB-14','Driftwood Studio','London','Paris','France','Business address','20-11-2026','Most reliable',1,20,4400,170,'No'),
    [it('TB-14','NAP-01','Cotton napkin set of 4',630251,400,11,20,20,0.3,'Vietnam')],
    'Nothing — business address, so no appointment needed','Green','— (R-08 not applicable)','Carrier C · Road express (£330)')
add(bk('TB-15','Quayside Gifts','Leeds','Berlin','Germany','Warehouse','01-10-2026','Cheapest',2,50,5000,400,'No'),
    [it('TB-15','MUG-01','Ceramic mug, 350 ml',None,360,6,12,30,0.45,'China'),
     it('TB-15','CDL-02','Scented candle, 200 g',340600,240,8.25,16,15,0.35,None)],
    'Several at once: no HS, value off, cartons off, no appointment, no origin, past date','Red','R-01, R-02, R-04 (Red) + R-08, R-09, R-10 (Amber)','Carrier A · Road economy (£360)')
add(bk('TB-16','Cobble Lane Pantry','Manchester','Rotterdam','Netherlands','Warehouse',datetime.datetime(2026,11,24),'Balanced',1,30,2800,300,'yes'),
    [it('TB-16','BWL-01','Ceramic serving bowl','6912.00',120,10,6,20,1.1,'UK'),
     it('TB-16','TEA-02','Black tea, loose leaf 250 g','090240',400,4,40,10,0.27,'India')],
    'Format edge cases only: HS “6912.00”, HS “090240” with leading zero as text, origin “UK”, date as a real Excel date, “yes” in lower case','Green','—','Carrier B · Road standard (£200)')
add(bk('TB-17','Pennine Pottery','Manchester','Frankfurt','Germany','Warehouse','18-11-2026','Cheapest',3,75,5454,1500,'Yes'),
    [it('TB-17','PLT-01','Stoneware dinner plate, 27 cm',691200,900,6,12,75,1.4,'United Kingdom')],
    'Boundaries: declared exactly +1.0% over lines; exactly 500 kg per pallet','Green','— (both limits are inclusive)','Carrier A · Road economy (£540)')
add(bk('TB-18','Atlas Prints','London','Paris','Germany','Warehouse','19-11-2026','Cheapest',1,10,2000,150,'Yes'),
    [it('TB-18','PRT-01','Framed art print, A3',491191,100,20,10,10,1.2,'United Kingdom')],
    'Destination city Paris does not match country Germany','Red','Data: city/country mismatch','Carrier A · Road economy (£170, France rates once fixed)')
add(bk('TB-19','Mossbank Ltd','Leeds','Amsterdam','Netherlands','Warehouse','18-11-2026','Cheapest',1,10,1000,100,'Yes'),
    [],
    'Booking has no item lines','Red','Data: no items (R-01–R-04, R-06, R-07, R-09 cannot run)','Carrier A · Road economy (£160)')
add(bk('TB-20','Hearth & Flue','Birmingham','Hamburg','Germany','Business address','23-11-2026','Cheapest',2,20,9000,1600,'No'),
    [it('TB-20','STV-01','Cast iron log burner stove',732189,20,450,1,20,75,'Turkey')],
    '800 kg per pallet — over every carrier limit (max 750)','Red','R-05 (fix: Carrier B on 3 pallets, £660)','Re-pack → Carrier B · Road standard (£660)')
add(bk('TB-21','Confetti & Co','London','Paris','France','Business address','19-11-2026','Cheapest',1,20,1600,150,'No'),
    [it('TB-21','PP-01','Party poppers, box of 50',950590,200,8,10,20,0.4,'China')],
    'Party poppers — ADR class 1 (fireworks), outside the R-06 categories','Green','— (Advisory: possible dangerous goods, ADR class 1)','Carrier A · Road economy (£170)')
add(bk('TB-22','Blue Lagoon Pools','Birmingham','Amsterdam','Netherlands','Business address','20-11-2026','Cheapest',2,25,3000,600,'No'),
    [it('TB-22','CHL-01','Pool chlorine granules 5 kg tub',380894,100,30,4,25,5.2,'China')],
    'Pool chlorine — ADR class 5.1 (oxidiser), outside the R-06 categories','Green','— (Advisory: possible dangerous goods, ADR class 5.1)','Carrier A · Road economy (£320)')
add(bk('TB-23','Toolbox Toys','Leeds','Berlin','Germany','Business address','18-11-2026','Cheapest',1,30,2400,200,'No'),
    [it('TB-23','SL-01','Spirit level 60 cm',901580,100,12,10,10,0.5,'China'),
     it('TB-23','MC-01','Matchbox car set',950300,200,6,10,20,0.3,'China')],
    'Look-alike words: “spirit level” and “matchbox car” are not dangerous goods','Green','—','Carrier A · Road economy (£180)')
add(bk('TB-24','Loom & Larder','London','Lyon','France','Business address','19-11-2026','Cheapest',4,40,4000,600,'No'),
    [it('TB-24','BLK-01','Wool blanket, double',630120,200,20,5,40,2.0,'United Kingdom')],
    '600 kg on 4 pallets — 2 pallets are enough by weight','Green','— (Advisory: 2 pallets enough, saving £340)','Carrier A · Road economy (£680)')
add(bk(' tb-05 ','Alder Home','London','Hamburg','Germany','Warehouse','23-11-2026','Most reliable',1,10,900,120,'Yes'),
    [],
    'Booking ID “ tb-05 ” repeats TB-05 (different case and spaces)','Red','Data: duplicate Booking ID (its items cannot be matched)','Carrier C · Road express (£340)')
items.append(['TB-99','GHOST-01','Item pointing at a booking that does not exist',691200,10,1,10,1,0.5,'China'])
expected.append(['TB-99 (Items only)','Item row whose Booking ID is not on the Bookings sheet','Notice','Reported as an orphan item row, not checked','—'])

def style_like(ws, src_row, dst_row, ncol):
    for c in range(1, ncol+1):
        s = ws.cell(src_row, c); d = ws.cell(dst_row, c)
        d.font = copy.copy(s.font); d.border = copy.copy(s.border); d.alignment = copy.copy(s.alignment)
for r, row in enumerate(bookings, start=4):
    for c, v in enumerate(row, start=1):
        cell = B.cell(r, c, v)
        if isinstance(v, datetime.datetime): cell.number_format = 'DD-MM-YYYY'
for r, row in enumerate(items, start=4):
    for c, v in enumerate(row, start=1):
        cell = I.cell(r, c, v)
        if c == 4 and isinstance(v, str): cell.number_format = '@'

# Expected results sheet
E = wb.create_sheet('Expected results', 2)
hdr = ['Booking ID', 'Planted mistake / edge case', 'Expected status', 'Rules expected to flag', 'Expected carrier']
E['A1'] = 'Test cases — what Preflight should report for each booking above (advisories never change the status) (checked on or before 16-11-2026; after that, every date becomes “past” and R-10 fires).'
E['A1'].font = Font(name='Arial', bold=True, size=11)
for c, h in enumerate(hdr, 1):
    x = E.cell(2, c, h); x.font = Font(name='Arial', bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F3864'); x.alignment = Alignment(vertical='center', wrap_text=True)
fills = {'Red': 'F8D7D3', 'Amber': 'FCEBC7', 'Green': 'D6EFDF', 'Notice': 'E7ECF1'}
thin = Side(style='thin', color='C9D2DA')
for r, row in enumerate(expected, 3):
    for c, v in enumerate(row, 1):
        x = E.cell(r, c, v); x.font = Font(name='Arial', size=10); x.alignment = Alignment(wrap_text=True, vertical='top'); x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    E.cell(r, 3).fill = PatternFill('solid', fgColor=fills[row[2]])
for c, w in enumerate([12, 52, 12, 40, 38], 1): E.column_dimensions[get_column_letter(c)].width = w
E.freeze_panes = 'A3'
wb.save('../assets/Preflight_Test_Cases.xlsx')
print(len(bookings), 'bookings,', len(items), 'item rows')
